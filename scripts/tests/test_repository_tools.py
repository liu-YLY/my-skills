import importlib.util
import json
from collections import Counter
from pathlib import Path

import pytest


def test_docx_conversion_preserves_table_section_order(tmp_path):
    from docx import Document

    source = tmp_path / "requirements.docx"
    document = Document()
    document.add_heading("支付规则", level=1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "退款条件"
    table.cell(0, 1).text = "未履约"
    document.add_heading("注册规则", level=1)
    document.add_paragraph("邮箱唯一")
    document.save(source)
    script = Path(__file__).resolve().parents[2] / "plugins/testing/skills/test-case-engineer/scripts/convert_docs.py"
    spec = importlib.util.spec_from_file_location("doc_converter", script)
    converter = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(converter)
    output = converter.convert_file(source)[0].read_text()
    assert output.index("# 支付规则") < output.index("| 退款条件") < output.index("# 注册规则")
    assert "邮箱唯一" in output


def test_docx_conversion_preserves_adjacent_tables_and_text(tmp_path):
    from docx import Document

    source = tmp_path / "requirements.docx"
    document = Document()
    document.add_paragraph("开始")
    for text in ("第一表|字段\n第二行", "第二表"):
        document.add_table(rows=1, cols=1).cell(0, 0).text = text
    document.add_paragraph("")
    document.add_heading("结束", level=2)
    document.save(source)
    script = Path(__file__).resolve().parents[2] / "plugins/testing/skills/test-case-engineer/scripts/convert_docs.py"
    spec = importlib.util.spec_from_file_location("doc_converter", script)
    converter = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(converter)
    output = converter.convert_file(source)[0].read_text()
    assert output.index("开始") < output.index("第一表") < output.index("第二表") < output.index("## 结束")
    assert "第一表\\|字段<br>第二行" in output

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("skill_evals", ROOT / "scripts/skill-evals.py")
evals = importlib.util.module_from_spec(spec)
spec.loader.exec_module(evals)
version_spec = importlib.util.spec_from_file_location(
    "version_sync", ROOT / "scripts/check-version-sync.py"
)
version_sync = importlib.util.module_from_spec(version_spec)
version_spec.loader.exec_module(version_sync)


@pytest.mark.parametrize("payload", [
    [{"id": 1, "prompt": "task", "expected": "observable result"}],
    {"evals": [{"id": 1, "prompt": "task", "expectations": ["observable result"]}]},
    {"prompts": [{"id": 1, "prompt": "task", "expected_behavior": "observable result",
                   "validation_points": ["evidence"]}]},
])
def test_eval_formats_preserve_prompt_and_expectations(tmp_path, payload):
    path = tmp_path / "test-prompts.json"
    path.write_text(json.dumps(payload))
    cases = evals.read_cases(path)
    assert cases[0]["prompt"] == "task"
    assert "observable result" in cases[0]["checks"]


def test_eval_reader_rejects_duplicate_ids(tmp_path):
    path = tmp_path / "test-prompts.json"
    sample = dict(id=1, prompt="task", expected="result")
    path.write_text(json.dumps([sample, sample]))
    with pytest.raises(ValueError, match="duplicate"):
        evals.read_cases(path)


def test_eval_reader_preserves_negative_trigger(tmp_path):
    path = tmp_path / "test-prompts.json"
    path.write_text(json.dumps({"trigger_evals": [{"query": "translate", "should_trigger": False}]}))
    assert evals.read_cases(path)[0]["should_trigger"] is False


def test_every_skill_has_complete_trigger_coverage():
    cases = evals.load_cases(ROOT)
    evals.validate_trigger_coverage(cases)
    trigger_counts = Counter(case["skill"] for case in cases if case["kind"] == "trigger")
    assert set(trigger_counts) == {
        "bug-analyzer",
        "change-impact-analyzer",
        "performance-test-engineer",
        "state-machine-test-engineer",
        "test-case-engineer",
        "test-data-generator",
        "test-strategy-engineer",
        "testing-bundle",
        "wechat-formatter",
    }
    assert min(trigger_counts.values()) >= 4


def test_root_readme_skill_versions_follow_frontmatter(tmp_path):
    skill_dir = tmp_path / "plugins/testing/skills/sample-skill"
    skill_dir.mkdir(parents=True)
    (skill_dir / "SKILL.md").write_text(
        "---\nname: sample-skill\nversion: 1.2.1\n---\n", encoding="utf-8"
    )
    (tmp_path / "README.md").write_text(
        "| Skill | 版本 |\n|---|---|\n| sample-skill | v1.2.0 |\n", encoding="utf-8"
    )
    errors = version_sync.check_root_readme_skill_versions(tmp_path)
    assert errors == [
        f"{tmp_path / 'README.md'}: references 'sample-skill v1.2.0' "
        "but SKILL.md frontmatter version is 1.2.1"
    ]


def test_current_changes_include_committed_feature_branch_work():
    skill_dir = ROOT / "plugins/testing/skills/change-impact-analyzer"
    skill = (skill_dir / "SKILL.md").read_text(encoding="utf-8")
    modes = (skill_dir / "knowledge/diff-modes.md").read_text(encoding="utf-8")
    assert "工作区干净且 feature 分支领先默认分支" in skill
    assert "不得忽略 feature 分支上已提交但未合并的改动" in modes
    assert "git diff --no-ext-diff --no-textconv <default>...HEAD" in modes


def test_bundle_declares_current_state_machine_mcp_version():
    bundle = ROOT / "plugins/testing/skills/testing-bundle"
    package = ROOT / "plugins/testing/mcp-servers/state-machine-testing/pyproject.toml"
    version_line = next(line for line in package.read_text(encoding="utf-8").splitlines()
                        if line.startswith("version = "))
    version = version_line.split('"')[1]
    for path in (bundle / "SKILL.md", bundle / "knowledge/usage-examples.md"):
        content = path.read_text(encoding="utf-8")
        assert f"v{version}" in content


def test_state_machine_runtime_docs_do_not_hardcode_test_totals():
    skill_dir = ROOT / "plugins/testing/skills/state-machine-test-engineer"
    for path in (skill_dir / "SKILL.md", skill_dir / "README.md",
                 skill_dir / "integrations/quickstart.md"):
        content = path.read_text(encoding="utf-8")
        assert "项测试全绿" not in content


def test_strategy_does_not_invent_missing_risk_scores_or_force_ratios():
    skill = (ROOT / "plugins/testing/skills/test-strategy-engineer/SKILL.md").read_text(
        encoding="utf-8"
    )
    assert "缺失维度标「待补」且不计算正式总分" in skill
    assert "不得为满足固定比例下调风险等级" in skill
    assert "P0 占比 ≤ 30%" not in skill
    assert "前 30% → P0" not in skill


def test_performance_diagnosis_is_evidence_first_when_data_is_missing():
    skill_dir = ROOT / "plugins/testing/skills/performance-test-engineer"
    skill = (skill_dir / "SKILL.md").read_text(encoding="utf-8")
    prompts = json.loads((skill_dir / "test-prompts.json").read_text(encoding="utf-8"))
    diagnosis = next(item for item in prompts["prompts"] if item["id"] == 2)["expected"]
    vague_request = next(item for item in prompts["prompts"] if item["id"] == 4)["expected"]
    assert "不按固定顺序确认根因" in diagnosis
    assert "按 应用层→资源层→架构层 顺序" not in diagnosis
    assert "标为待定" in vague_request
    assert "默认值兜底" not in vague_request
    assert "缺 TPS 时不得推断吞吐拐点或资源饱和" in skill
    assert "不得输出已确认瓶颈" in skill


def test_case_review_respects_existing_revision_authority_and_risk_evidence():
    skill_dir = ROOT / "plugins/testing/skills/test-case-engineer"
    review = (skill_dir / "knowledge/review-mode.md").read_text(encoding="utf-8")
    core = (skill_dir / "test-case-engineer-core.md").read_text(encoding="utf-8")
    prompts = json.loads((skill_dir / "test-prompts.json").read_text(encoding="utf-8"))
    revision = next(item for item in prompts["evals"] if item["id"] == 14)["expectations"]
    priority = next(item for item in prompts["evals"] if item["id"] == 15)["expectations"]
    assert any("不重复要求确认" in item for item in revision)
    assert any("不得仅凭" in item and "固定占比" in item for item in priority)
    assert "P0 占比应 < 30%" not in review
    assert "不重跑与当前问题无关的完整四阶段" in review
    assert "优先级比例严重偏离" not in core


def write_model_output_fixtures(path, rows):
    for row in rows:
        if row.get("output_path"):
            output = path.parent / row["output_path"]
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text("x" * row["output_characters"])


def test_model_results_are_validated_and_summarized(tmp_path):
    cases = [
        dict(skill="sample", id="task-1", kind="task", prompt="do it",
             checks=["complete", "grounded"], should_trigger=None),
        dict(skill="sample", id="trigger-1", kind="trigger", prompt="route it",
             checks=[], should_trigger=True),
        dict(skill="sample", id="trigger-2", kind="trigger", prompt="translate it",
             checks=[], should_trigger=False),
    ]
    rows = [
        dict(skill="sample", id="task-1", kind="task", variant="with_skill", run=1,
             duration_seconds=2.0, output_characters=100, output_path="outputs/task-1.md",
             check_results=[True, False]),
        dict(skill="sample", id="trigger-1", kind="trigger", variant="with_skill", run=1,
             duration_seconds=1.0, output_characters=20,
             output_path="outputs/trigger-1.md", predicted_trigger=True),
        dict(skill="sample", id="trigger-2", kind="trigger", variant="with_skill", run=1,
             duration_seconds=1.0, output_characters=20,
             output_path="outputs/trigger-2.md", predicted_trigger=False),
    ]
    path = tmp_path / "results.jsonl"
    path.write_text("\n".join(json.dumps(row) for row in rows), encoding="utf-8")
    write_model_output_fixtures(path, rows)
    results = evals.read_results(path, cases)
    result_summary = evals.summarize_results(results, cases)
    summary = result_summary["variants"]["with_skill"]
    assert summary["checks_passed"] == summary["checks_total"] == 0
    assert summary["trigger_precision"] is summary["trigger_recall"] is None
    assert summary["mean_duration_seconds"] is None
    assert len(summary["unpaired_runs"]) == 3
    assert summary["completion_rate"] == 1.0
    assert summary["missing_cases"] == []
    baseline = evals.summarize_results(results, cases)["variants"]["baseline"]
    assert baseline["completion_rate"] == 0.0
    assert len(baseline["missing_cases"]) == 3
    assert result_summary["status"] == "partial"


def test_model_results_report_partial_case_completion(tmp_path):
    cases = [
        dict(skill="sample", id="task-1", kind="task", prompt="do it",
             checks=["complete"], should_trigger=None),
        dict(skill="sample", id="task-2", kind="task", prompt="do more",
             checks=["grounded"], should_trigger=None),
    ]
    row = dict(skill="sample", id="task-1", kind="task", variant="with_skill", run=1,
               duration_seconds=1.0, output_characters=10, output_path="outputs/task-1.md",
               check_results=[True])
    path = tmp_path / "results.jsonl"
    path.write_text(json.dumps(row), encoding="utf-8")
    write_model_output_fixtures(path, [row])
    summary = evals.summarize_results(evals.read_results(path, cases), cases)
    variant = summary["variants"]["with_skill"]
    assert variant["completed_cases"] == 1
    assert variant["expected_cases"] == 2
    assert variant["completion_rate"] == 0.5
    assert variant["missing_cases"] == ["sample:task:task-2"]


def test_paired_evaluation_records_expose_evidence_gaps():
    path = ROOT / "docs/skill-evaluation/2026-09-21/paired-results.jsonl"
    rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
    required = {
        "evaluation_id", "skill", "baseline_ref", "candidate_ref", "prompt_ids",
        "judge_id", "verdict", "margin", "reason", "regression_risk",
        "duration_seconds", "output_path", "evidence_status", "evidence_gaps",
    }
    assert len(rows) == 6
    assert all(required <= row.keys() for row in rows)
    assert all(row["verdict"] in {"baseline", "candidate", "tie"} for row in rows)
    assert all(row["margin"] in {"slight", "clear", "large"} for row in rows)
    assert all(row["evidence_status"] in {"full", "partial"} for row in rows)
    for row in rows:
        if row["evidence_status"] == "full":
            assert isinstance(row["duration_seconds"], (int, float))
            assert row["duration_seconds"] >= 0
            assert row["regression_risk"].strip()
        else:
            assert row["evidence_gaps"]
    votes = Counter((row["evaluation_id"], row["verdict"]) for row in rows)
    assert votes[("performance-structure-20260921", "candidate")] == 3
    assert votes[("test-case-contract-20260921", "candidate")] == 3


@pytest.mark.parametrize("field,value,message", [
    ("run", 0, "positive integer"),
    ("variant", "unknown", "invalid variant"),
    ("check_results", [True], "2 booleans"),
    ("output_path", "", "output_path"),
])
def test_model_results_reject_invalid_contract(tmp_path, field, value, message):
    cases = [dict(skill="sample", id="task-1", kind="task", prompt="do it",
                  checks=["complete", "grounded"], should_trigger=None)]
    row = dict(skill="sample", id="task-1", kind="task", variant="with_skill", run=1,
               duration_seconds=2.0, output_characters=100, output_path="outputs/task-1.md",
               check_results=[True, True])
    row[field] = value
    path = tmp_path / "results.jsonl"
    path.write_text(json.dumps(row), encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        evals.read_results(path, cases)


@pytest.mark.parametrize("defect", ["missing", "hash", "length", "nan", "infinity"])
def test_model_artifact_failures_are_partial_and_excluded(tmp_path, defect):
    import hashlib

    output = tmp_path / "output.md"
    output.write_text("report")
    case = dict(skill="sample", id="1", kind="task", prompt="task", checks=["done"])
    row = dict(skill="sample", id="1", kind="task", run=1, duration_seconds=1.0,
               output_characters=6, output_path="output.md", check_results=[True],
               output_sha256=hashlib.sha256(output.read_bytes()).hexdigest())
    if defect == "missing":
        output.unlink()
    elif defect == "hash":
        row["output_sha256"] = "0" * 64
    elif defect == "length":
        row["output_characters"] = 7
    else:
        row["duration_seconds"] = float("nan" if defect == "nan" else "inf")
    path = tmp_path / "results.jsonl"
    path.write_text("\n".join(json.dumps(dict(row, variant=variant))
                              for variant in ("with_skill", "baseline")))
    summary = evals.summarize_results(evals.read_results(path, [case]), [case])
    assert summary["status"] == "partial"
    assert summary["variants"]["with_skill"]["checks_total"] == 0
    assert summary["variants"]["with_skill"]["mean_duration_seconds"] is None


def test_model_repetitions_require_each_expected_run_and_pair(tmp_path):
    case = dict(skill="sample", id="1", kind="task", prompt="task", checks=["done"])
    rows = [dict(skill="sample", id="1", kind="task", variant=variant, run=run,
                 duration_seconds=1.0, output_characters=4, output_path=f"{variant}-{run}.md",
                 check_results=[True])
            for variant in ("with_skill", "baseline") for run in (1, 2)]
    rows.append(dict(rows[0], run=3, output_path="candidate-3.md"))
    path = tmp_path / "results.jsonl"
    path.write_text("\n".join(json.dumps(row) for row in rows))
    write_model_output_fixtures(path, rows)
    summary = evals.summarize_results(evals.read_results(path, [case]), [case], expected_runs=3)
    assert summary["status"] == "partial"
    baseline = summary["variants"]["baseline"]
    candidate = summary["variants"]["with_skill"]
    assert baseline["completed_cases"] == 0 and baseline["completed_runs"] == 2
    assert candidate["completed_cases"] == 1 and candidate["completed_runs"] == 3
    assert candidate["checks_total"] == baseline["checks_total"] == 2
    assert candidate["unpaired_runs"] == ["sample:task:1:run-3"]
    assert baseline["missing_runs"] == ["sample:task:1:run-3"]
    rows.append(dict(rows[-1], variant="baseline", output_path="baseline-3.md"))
    write_model_output_fixtures(path, rows)
    path.write_text("\n".join(json.dumps(row) for row in rows))
    complete = evals.summarize_results(evals.read_results(path, [case]), [case], expected_runs=3)
    assert complete["status"] == "evaluated"
    assert complete["variants"]["baseline"]["checks_total"] == 3


def test_selected_frozen_suite_is_independent_from_repository_cases(tmp_path):
    import subprocess
    import sys

    path = tmp_path / "frozen.json"
    path.write_text(json.dumps({"task_cases": [{"id": "one", "skill": "sample", "prompt": "task",
                                              "expectations": ["done"]}],
                               "original_trigger_cases": [{"id": "route", "skill": "sample",
                                  "original": {"query": "route", "should_trigger": False,
                                               "category": "negative"}}]}))
    result = subprocess.run([sys.executable, str(ROOT / "scripts/skill-evals.py"),
                             "--cases-file", str(path), "--json"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    cases = json.loads(result.stdout)
    assert len(cases) == 2 and cases[0]["id"] == "one"
    assert cases[1]["should_trigger"] is False


def test_model_pair_metrics_exclude_a_valid_but_unpaired_artifact(tmp_path):
    case = dict(skill="sample", id="1", kind="task", prompt="task", checks=["done"])
    rows = [dict(skill="sample", id="1", kind="task", variant=variant, run=1,
                 duration_seconds=1.0, output_characters=4, output_path=f"{variant}.md",
                 check_results=[True]) for variant in ("with_skill", "baseline")]
    path = tmp_path / "results.jsonl"
    write_model_output_fixtures(path, rows)
    (tmp_path / "baseline.md").unlink()
    path.write_text("\n".join(json.dumps(row) for row in rows))
    summary = evals.summarize_results(evals.read_results(path, [case]), [case])
    assert summary["status"] == "partial"
    assert summary["variants"]["with_skill"]["completed_runs"] == 1
    assert summary["variants"]["with_skill"]["checks_total"] == 0
    assert summary["variants"]["baseline"]["invalid_runs"]


def test_model_trigger_repeats_can_differ_from_task_repeats(tmp_path):
    cases = [dict(skill="sample", id="task", kind="task", prompt="task", checks=["done"]),
             dict(skill="sample", id="trigger", kind="trigger", prompt="route", checks=[], should_trigger=True)]
    rows=[]
    for variant in ("old", "new", "no_skill"):
        for case in cases:
            for run in range(1, 4 if case["kind"] == "task" else 2):
                row=dict(skill="sample", id=case["id"], kind=case["kind"], variant=variant,
                         run=run, duration_seconds=1.0, output_characters=4,
                         output_path=f"{variant}-{case['id']}-{run}.md")
                row.update(check_results=[True]) if case["kind"] == "task" else row.update(predicted_trigger=True)
                rows.append(row)
    path=tmp_path / "results.jsonl"
    write_model_output_fixtures(path, rows)
    path.write_text("\n".join(json.dumps(row) for row in rows))
    read=evals.read_results(path, cases)
    summary=evals.summarize_results(read, cases, expected_runs=3, expected_trigger_runs=1)
    assert summary["status"] == "evaluated"
    assert summary["variants"]["old"]["expected_runs"] == 4
    assert summary["variants"]["old"]["trigger_runs"] == 1
    assert evals.summarize_results(read, cases, expected_runs=3)["status"] == "partial"
